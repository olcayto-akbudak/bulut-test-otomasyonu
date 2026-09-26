"""
Ortak ayarlar:
  1) .env'den ortam bilgisi okur
  2) Canlı ortam koruması (yanlış adrese test koşturmayı engeller)
  3) Oturumu bir kez açıp tüm testlerde yeniden kullanır (storage_state)
  4) Koşu sonunda Excel raporu + TaskPano'ya yapıştırılacak metin üretir
"""
import os
from datetime import datetime
from pathlib import Path

import pytest
from dotenv import load_dotenv

load_dotenv()

KOK = Path(__file__).parent
AUTH_DOSYASI = KOK / ".auth" / "oturum.json"
RAPOR_KLASORU = KOK / "raporlar"


# ---------------------------------------------------------------- ortam
def _env(ad: str, varsayilan: str = "") -> str:
    return os.getenv(ad, varsayilan).strip()


def pytest_sessionstart(session):
    """Canlı ortam koruması: tek bir test bile başlamadan önce çalışır."""
    base_url = _env("BASE_URL")
    if not base_url:
        pytest.exit("BASE_URL tanımlı değil. .env.example dosyasını .env olarak kopyalayıp doldur.", 2)
    yasaklar = [y.strip() for y in _env("YASAK_ADRESLER", "app.,www.,canli,prod").split(",") if y.strip()]
    for y in yasaklar:
        if y in base_url.lower():
            pytest.exit(f"GÜVENLİK: BASE_URL '{y}' içeriyor, canlı ortam olabilir. Testler durduruldu.", 3)


@pytest.fixture(scope="session")
def ortam():
    base_url = _env("BASE_URL")
    return {
        "base_url": base_url.rstrip("/"),
        "kullanici": _env("TEST_USER"),
        "sifre": _env("TEST_PASS"),
    }


@pytest.fixture(scope="session")
def browser_context_args(browser_context_args, ortam):
    # Tüm sayfalarda page.goto("/cari") gibi göreli adres kullanabilmek için
    return {**browser_context_args, "base_url": ortam["base_url"], "locale": "tr-TR"}


# ---------------------------------------------------------------- oturum
@pytest.fixture(scope="session")
def oturum_dosyasi(browser, browser_context_args, ortam):
    """Girişi koşu başında BİR KEZ yapar; sonraki testler bu oturumu kullanır."""
    from pages.login_page import LoginPage

    AUTH_DOSYASI.parent.mkdir(exist_ok=True)
    context = browser.new_context(**browser_context_args)
    page = context.new_page()
    LoginPage(page).ac().giris_yap(ortam["kullanici"], ortam["sifre"])
    LoginPage(page).giris_basarili_mi()
    context.storage_state(path=str(AUTH_DOSYASI))
    context.close()
    return str(AUTH_DOSYASI)


@pytest.fixture
def sayfa(new_context, oturum_dosyasi):
    """Giriş yapılmış hazır sayfa. Testlerin çoğu bunu kullanır.
    new_context: hata anında ekran görüntüsü, video ve trace'i otomatik kaydeder."""
    return new_context(storage_state=oturum_dosyasi).new_page()


# ---------------------------------------------------------------- raporlama
def pytest_configure(config):
    config.addinivalue_line("markers", "senaryo(id): Test Senaryo Yönetimi aracındaki senaryo numarası")
    config._sonuclar = []


def pytest_runtest_logreport(report):
    # setup hatası (ör. giriş başarısız) veya asıl test adımı
    if report.when == "call" or (report.when == "setup" and report.outcome != "passed"):
        senaryo = next((m.args[0] for m in report.item_markers if m.name == "senaryo"), "") \
            if hasattr(report, "item_markers") else ""
        hata = ""
        if report.failed:
            # pytest'in "E   ..." satırları asıl hata mesajıdır (beklenen / gerçekleşen)
            e_satirlari = [l[1:].strip() for l in report.longreprtext.splitlines() if l.startswith("E ")]
            hata = (" | ".join(e_satirlari[:2]) or report.longreprtext.strip().splitlines()[-1])[:300]
        elif report.skipped:
            hata = str(report.longrepr[-1]) if isinstance(report.longrepr, tuple) else ""
        report.config_ref._sonuclar.append({
            "senaryo": senaryo,
            "test": report.nodeid.split("::")[-1],
            "dosya": report.nodeid.split("::")[0],
            "sonuc": {"passed": "GEÇTİ", "failed": "KALDI", "skipped": "ATLANDI"}[report.outcome],
            "sure": round(report.duration, 2),
            "hata": hata,
        })


@pytest.hookimpl(hookwrapper=True)
def pytest_runtest_makereport(item, call):
    outcome = yield
    rep = outcome.get_result()
    rep.item_markers = list(item.iter_markers())
    rep.config_ref = item.config


def pytest_sessionfinish(session, exitstatus):
    sonuclar = getattr(session.config, "_sonuclar", [])
    if not sonuclar:
        return
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    RAPOR_KLASORU.mkdir(exist_ok=True)
    zaman = datetime.now()
    meta = {
        "Tarih": zaman.strftime("%d.%m.%Y %H:%M"),
        "Tester": _env("TESTER"),
        "Program": _env("PROGRAM"),
        "Sürüm": _env("SURUM"),
        "İş Kaydı No": _env("IS_KAYDI_NO"),
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Sonuçları"
    basliklar = list(meta.keys()) + ["Senaryo", "Test", "Dosya", "Sonuç", "Süre (sn)", "Hata Özeti"]
    ws.append(basliklar)
    for h in ws[1]:
        h.font = Font(bold=True)
    renk = {"GEÇTİ": "C6EFCE", "KALDI": "FFC7CE", "ATLANDI": "FFEB9C"}
    for s in sonuclar:
        ws.append(list(meta.values()) + [s["senaryo"], s["test"], s["dosya"], s["sonuc"], s["sure"], s["hata"]])
        ws.cell(ws.max_row, len(meta) + 4).fill = PatternFill("solid", fgColor=renk[s["sonuc"]])
    for kolon in ws.columns:
        ws.column_dimensions[kolon[0].column_letter].width = min(60, max(len(str(c.value or "")) for c in kolon) + 2)

    ad = f"test_raporu_{zaman:%Y%m%d_%H%M}"
    wb.save(RAPOR_KLASORU / f"{ad}.xlsx")

    # TaskPano'ya yapıştırılacak kısa özet
    gecen = sum(s["sonuc"] == "GEÇTİ" for s in sonuclar)
    kalan = [s for s in sonuclar if s["sonuc"] == "KALDI"]
    satirlar = [
        f"[{meta['Program']} {meta['Sürüm']}] Otomasyon testi — {meta['Tarih']} — {meta['Tester']}",
        f"İş Kaydı: {meta['İş Kaydı No'] or '-'} | Toplam: {len(sonuclar)} | Geçti: {gecen} | Kaldı: {len(kalan)}",
    ]
    for s in kalan:
        satirlar.append(f"  ✗ {s['senaryo'] or s['test']}: {s['hata']}")
    (RAPOR_KLASORU / f"{ad}_taskpano.txt").write_text("\n".join(satirlar), encoding="utf-8")
